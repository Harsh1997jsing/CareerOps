"""
Shared DOCX assembly for generated documents. Every document built here is
single-column, headings-and-paragraphs only — no tables, text boxes, or
images — since those are the layout elements that most often break ATS
text extraction. ats_validator.py (Phase 3) checks documents stay this way.
"""

from docx import Document

from app.llm.schemas import GeneratedResumeSection

SECTION_TITLES = {
    "summary": "Summary",
    "skills": "Skills",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
}


def write_resume_docx(profile: dict, sections: list[GeneratedResumeSection], output_path: str) -> None:
    """Assemble and save a clean, single-column, ATS-compatible resume in DOCX format.

    Adds candidate name heading, pipe-separated contact info line, and section
    headings with associated content paragraphs. Strictly avoids tables, multi-columns,
    text boxes, and inline shapes to guarantee maximum ATS readability.

    Args:
        profile: Candidate profile dict containing 'name', 'email', 'phone', etc.
        sections: Generated resume sections with section keys and formatted text.
        output_path: Destination filepath to write the `.docx` document to.
    """
    document = Document()

    document.add_heading(profile.get("name", ""), level=0)
    contact_line = " | ".join(filter(None, [
        profile.get("email"),
        profile.get("phone"),
        profile.get("location"),
        profile.get("linkedin"),
        profile.get("github"),
    ]))
    if contact_line:
        document.add_paragraph(contact_line)

    for section in sections:
        document.add_heading(SECTION_TITLES.get(section.section, section.section.title()), level=1)
        document.add_paragraph(section.content)

    document.save(output_path)


def write_cover_letter_docx(profile: dict, content: str, output_path: str) -> None:
    """Assemble and save an ATS-friendly cover letter in DOCX format.

    Adds candidate name, contact line, and body paragraphs separated by empty lines.
    Ensures layout is strictly linear and ATS-compliant.

    Args:
        profile: Candidate profile dict containing 'name', 'email', and 'phone'.
        content: Full text of the generated cover letter.
        output_path: Destination filepath to write the `.docx` file to.
    """
    document = Document()

    document.add_paragraph(profile.get("name", ""))
    contact_line = " | ".join(filter(None, [profile.get("email"), profile.get("phone")]))
    if contact_line:
        document.add_paragraph(contact_line)
    document.add_paragraph("")

    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)

    document.save(output_path)
