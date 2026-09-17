from docx import Document

from app.llm.schemas import GeneratedResumeSection
from app.services.docx_writer import write_cover_letter_docx, write_resume_docx

PROFILE = {
    "name": "Test Candidate",
    "email": "test@example.com",
    "phone": "+91-0000000000",
    "location": "Remote",
    "linkedin": "https://linkedin.com/in/test",
    "github": "https://github.com/test",
}


def test_write_resume_docx_contains_sections_and_no_illegal_elements(tmp_path):
    sections = [
        GeneratedResumeSection(section="summary", content="A concise summary.", evidence_ids_used=["EXP001"]),
        GeneratedResumeSection(section="skills", content="Python, FastAPI.", evidence_ids_used=["EXP001"]),
    ]
    output_path = tmp_path / "resume.docx"

    write_resume_docx(PROFILE, sections, str(output_path))

    assert output_path.exists()
    document = Document(str(output_path))

    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 0

    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "Test Candidate" in all_text
    assert "test@example.com" in all_text
    assert "A concise summary." in all_text
    assert "Python, FastAPI." in all_text
    assert "Summary" in all_text
    assert "Skills" in all_text


def test_write_cover_letter_docx_contains_body_and_no_illegal_elements(tmp_path):
    content = "Dear Hiring Manager,\n\nI am excited to apply.\n\nSincerely,\nTest Candidate"
    output_path = tmp_path / "cover_letter.docx"

    write_cover_letter_docx(PROFILE, content, str(output_path))

    assert output_path.exists()
    document = Document(str(output_path))

    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 0

    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "Dear Hiring Manager," in all_text
    assert "I am excited to apply." in all_text
    assert "Test Candidate" in all_text
