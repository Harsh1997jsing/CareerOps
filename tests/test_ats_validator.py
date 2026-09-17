import shutil
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.llm.schemas import GeneratedResumeSection
from app.services.ats_validator import (
    check_docx_structure,
    check_text_roundtrip,
    convert_docx_to_pdf,
    cover_letter_required_snippets,
    resume_required_snippets,
)
from app.services.docx_writer import write_resume_docx

PROFILE = {"name": "Test Candidate", "email": "test@example.com"}
SECTIONS = [
    GeneratedResumeSection(section="summary", content="A summary.", evidence_ids_used=["EXP001"]),
    GeneratedResumeSection(section="skills", content="Python, FastAPI.", evidence_ids_used=["EXP001"]),
]


def test_check_docx_structure_passes_for_clean_document(tmp_path):
    docx_path = tmp_path / "resume.docx"
    write_resume_docx(PROFILE, SECTIONS, str(docx_path))

    result = check_docx_structure(str(docx_path))

    assert result.passed
    assert result.reasons == []


def test_check_docx_structure_flags_table(tmp_path):
    docx_path = tmp_path / "with_table.docx"
    document = Document()
    document.add_paragraph("Hello")
    document.add_table(rows=1, cols=2)
    document.save(str(docx_path))

    result = check_docx_structure(str(docx_path))

    assert not result.passed
    assert any("table" in r for r in result.reasons)


def test_check_docx_structure_flags_multi_column_section(tmp_path):
    docx_path = tmp_path / "multi_column.docx"
    document = Document()
    document.add_paragraph("Hello")
    sect_pr = document.sections[0]._sectPr
    # Word represents column count as w:num on the section's single w:cols
    # element (already present by default) — not a second, separate element.
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = sect_pr.makeelement(qn("w:cols"), {})
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    document.save(str(docx_path))

    result = check_docx_structure(str(docx_path))

    assert not result.passed
    assert any("column" in r for r in result.reasons)


def test_check_text_roundtrip_passes_when_snippets_present():
    source = "Test Candidate test@example.com Summary A summary of experience."
    extracted = source
    result = check_text_roundtrip(source, extracted, ["Test Candidate", "Summary"])

    assert result.passed


def test_check_text_roundtrip_flags_missing_snippet():
    source = "Test Candidate test@example.com Summary text here."
    extracted = "Test Candidate Summary text here."  # email dropped
    result = check_text_roundtrip(source, extracted, ["Test Candidate", "test@example.com"])

    assert not result.passed
    assert any("test@example.com" in r for r in result.reasons)


def test_check_text_roundtrip_flags_heavy_word_loss():
    source = " ".join(f"word{i}" for i in range(100))
    extracted = " ".join(f"word{i}" for i in range(50))  # half the words missing

    result = check_text_roundtrip(source, extracted, [])

    assert not result.passed
    assert any("missing from extracted" in r for r in result.reasons)


def test_convert_docx_to_pdf_raises_without_libreoffice(tmp_path):
    with patch("app.services.ats_validator.shutil.which", return_value=None):
        with pytest.raises(RuntimeError, match="LibreOffice"):
            convert_docx_to_pdf(str(tmp_path / "resume.docx"), str(tmp_path))


def test_resume_required_snippets_includes_name_email_and_headings():
    snippets = resume_required_snippets(PROFILE, SECTIONS)

    assert "Test Candidate" in snippets
    assert "test@example.com" in snippets
    assert "Summary" in snippets
    assert "Skills" in snippets


def test_cover_letter_required_snippets_includes_name_and_email():
    snippets = cover_letter_required_snippets(PROFILE)

    assert snippets == ["Test Candidate", "test@example.com"]


@pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")),
    reason="LibreOffice not installed in this environment",
)
def test_full_ats_roundtrip_with_real_libreoffice(tmp_path):
    from app.services.ats_validator import validate_ats

    docx_path = tmp_path / "resume.docx"
    write_resume_docx(PROFILE, SECTIONS, str(docx_path))
    source_text = "Test Candidate test@example.com Summary A summary. Skills Python, FastAPI."

    result = validate_ats(
        str(docx_path), source_text,
        resume_required_snippets(PROFILE, SECTIONS), str(tmp_path),
    )

    assert result.passed
